import argparse
import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def text(value):
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def number(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return 0.0


def main():
    parser = argparse.ArgumentParser(
        description="Export achievements and run groupings from the EU4 workbook."
    )
    parser.add_argument(
        "--input", type=Path, default=ROOT / "EU4_Achievements.xlsx"
    )
    parser.add_argument(
        "--output", type=Path, default=ROOT / "data" / "achievements.json"
    )
    args = parser.parse_args()

    workbook = load_workbook(args.input, data_only=True, read_only=True)
    list_sheet = workbook["List"]
    list_headers = next(list_sheet.iter_rows(values_only=True))
    achievements = []
    achievements_by_name = {}

    for sheet_row, values in enumerate(
        list_sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        record = dict(zip(list_headers, values))
        name = text(record.get("Name"))
        if not name:
            continue
        achievement = {
            "id": f"achievement-{sheet_row}",
            "name": name,
            "description": text(record.get("Description")),
            "difficulty": text(record.get("Difficulty")) or "Unknown",
            "rarity": number(record.get("Percentage")),
            "points": number(record.get("Points")),
            "done": text(record.get("Done?")).upper() == "Y",
            "startingCountry": text(record.get("Started with")),
            "startingConditions": text(record.get("Starting conditions")),
            "requirements": text(record.get("Completion requirements")),
            "notes": "\n\n".join(
                value
                for value in (text(record.get("Notes1")), text(record.get("Notes2")))
                if value
            ),
        }
        achievements.append(achievement)
        achievements_by_name[name] = achievement

    runs_sheet = workbook["Runs"]
    run_headers = next(runs_sheet.iter_rows(values_only=True))
    grouped_runs = OrderedDict()
    for values in runs_sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(run_headers, values))
        run_id = text(record.get("Run ID"))
        name = text(record.get("Name"))
        if not run_id or run_id == "000" or not name or name.startswith("#"):
            continue
        run = grouped_runs.setdefault(
            run_id,
            {
                "id": run_id,
                "country": text(record.get("Starting country")),
                "difficulty": text(record.get("Run difficulty")) or "Unknown",
                "name": name,
                "customNation": text(
                    record.get("Cheesable with custom nation?")
                ),
                "achievementIds": [],
            },
        )
        achievement = achievements_by_name.get(name)
        if achievement is None:
            raise ValueError(
                f"Run {run_id} references achievement {name!r}, "
                "which is missing from the List worksheet."
            )
        if achievement["id"] not in run["achievementIds"]:
            run["achievementIds"].append(achievement["id"])

    data = {
        "source": args.input.name,
        "achievements": achievements,
        "runs": list(grouped_runs.values()),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        f"Exported {len(achievements)} achievements and {len(grouped_runs)} runs "
        f"to {args.output}"
    )


if __name__ == "__main__":
    main()
